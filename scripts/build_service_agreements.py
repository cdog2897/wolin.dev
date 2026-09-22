from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "agreements"

NAVY = "17324D"
SLATE = "435466"
MUTED = "66727E"
LIGHT = "E8EDF2"
BLACK = "000000"


PACKAGES = [
    {
        "name": "Starter",
        "price": "$279",
        "scope": "Google Business Profile",
        "description": "Hands-on care for the profile customers see in Google Search and Maps.",
        "included": [
            "A complete Google Business Profile audit with prioritized findings",
            "Updates to categories, services, description, attributes, and links where access and platform rules permit",
            "A check of hours, service areas, address, phone number, and contact details",
            "A review of the logo, cover image, and placement of existing photos",
            "One Google Business Profile post per monthly service period",
            "Identification of duplicate, suspended, or inaccurate listing issues",
            "A monthly report covering profile status, local ranking observations, completed work, and recommended next steps",
        ],
        "excluded": [
            "Website design, website development, website edits, technical SEO, or Search Console work",
            "Paid advertising, ad spend, boosted posts, or management of advertising accounts",
            "Review generation campaigns, review removal, or a promise that a platform will publish or remove any review",
            "Directory citation cleanup, social media management, photography, video production, or custom graphic design",
            "Third-party fees, paid software, domain fees, hosting fees, or platform charges",
        ],
        "specific_limits": [
            "Google controls approval, verification, suspension, reinstatement, display, and ranking of Business Profiles.",
            "Provider may identify or submit corrections for listing problems, but cannot guarantee that Google will accept, publish, verify, reinstate, or retain any change.",
        ],
    },
    {
        "name": "Growth",
        "price": "$679",
        "scope": "Google profile and website essentials",
        "description": "Profile management plus practical fixes that make the website easier to find and use.",
        "included": [
            "All services included in the Starter package",
            "Two Google Business Profile posts per monthly service period instead of one",
            "A website health review focused on the pages that matter most to the Client's services and locations",
            "Corrections to broken links, missing metadata, and heading issues on prioritized pages",
            "Basic on-page SEO for prioritized core service and location pages",
            "Improvements to page titles, descriptions, and local keyword copy",
            "Practical improvements to mobile layout, navigation, tap targets, and forms",
            "Practical image sizing, loading, and speed improvements",
            "Improvements to calls to action, click-to-call links, and contact paths",
            "A check of local business structured data and business name, address, and phone consistency",
            "Google Search Console setup and monitoring of identified issues",
        ],
        "excluded": [
            "A full website redesign, rebuild, replatforming, hosting migration, or custom web application",
            "Unlimited pages, unlimited revisions, ecommerce development, booking-system development, or custom integrations",
            "Premium themes, plugins, software licenses, stock media, hosting, domains, or other third-party fees",
            "Comprehensive technical SEO, large-scale content production, ongoing citation campaigns, or AI visibility monitoring",
            "Paid advertising, social media management, professional photography, video production, legal review, or security certification",
        ],
        "specific_limits": [
            "Website work is prioritized by likely impact and available monthly capacity. Not every identified issue or page will necessarily be completed in one service period.",
            "Some improvements require the Client's host, developer, theme vendor, plugin vendor, or another third party. Provider is not responsible for delays or limitations outside Provider's control.",
        ],
    },
    {
        "name": "Authority",
        "price": "$1,279",
        "scope": "Full search website and AI visibility",
        "description": "A broad monthly program for strengthening the business's local search presence, website, and visibility in AI-generated answers.",
        "included": [
            "All services included in the Growth package",
            "A full website and technical SEO audit with a prioritized roadmap",
            "Prioritized work on crawl, indexing, sitemap, robots, canonical, and redirect issues",
            "Improvements to site structure, internal linking, and structured data",
            "Work on Core Web Vitals and other high-impact performance issues",
            "Keyword, competitor, service, and location opportunity research",
            "Writing or improvement of prioritized service, location, FAQ, and authority content",
            "A review of conversion tracking, calls to action, and lead paths",
            "An AI visibility audit across ChatGPT, Claude, and Google AI experiences available to Provider",
            "Testing of representative customer prompts for mentions, factual accuracy, and cited sources",
            "Ongoing search and AI visibility monitoring with roadmap updates",
        ],
        "excluded": [
            "A full website redesign or rebuild, custom software, custom applications, or extensive ecommerce development",
            "Unlimited pages, unlimited content, unlimited revisions, or completion of every roadmap item in a single service period",
            "Paid advertising, ad spend, ongoing social media management, professional photography, or video production",
            "Premium software, data providers, plugins, themes, stock media, hosting, domains, or other third-party fees",
            "Legal, tax, accessibility, cybersecurity, public-relations, or reputation-management certification or advice",
        ],
        "specific_limits": [
            "Search audits and AI prompt tests are snapshots. Search results and AI answers can vary by user, location, device, account history, model, prompt wording, and date.",
            "Provider may recommend and implement prioritized improvements, but cannot control whether a search engine or AI system crawls, indexes, ranks, mentions, cites, or accurately describes the Client.",
        ],
    },
]


def set_cellless_bottom_border(paragraph, color=LIGHT, size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def set_run_font(run, name="Aptos", size=10.5, bold=False, color=BLACK, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_field(run, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_document(doc, package_name):
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.different_first_page_header_footer = False
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.widow_control = True

    title = doc.styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.keep_with_next = True
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size, before, after in [
        ("Heading 1", 14, 14, 5),
        ("Heading 2", 11, 9, 3),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    list_style = doc.styles["List Bullet"]
    list_style.font.name = "Aptos"
    list_style.font.size = Pt(10.25)
    list_style.paragraph_format.left_indent = Inches(0.24)
    list_style.paragraph_format.first_line_indent = Inches(-0.16)
    list_style.paragraph_format.space_after = Pt(3)
    list_style.paragraph_format.line_spacing = 1.06

    doc.core_properties.title = f"{package_name} Monthly Services Agreement"
    doc.core_properties.subject = "Month to month local visibility services agreement"
    doc.core_properties.author = "Caleb Wolin"
    doc.core_properties.keywords = "Wolin, services agreement, local search, monthly services"


def add_label_line(doc, label, blank="____________________________________________________________"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(f"{label}  ")
    set_run_font(r, size=10, bold=True, color=SLATE)
    r = p.add_run(blank)
    set_run_font(r, size=10, color=BLACK)
    return p


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_paragraph(doc, text=None, bold_lead=None, italic=False):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
    if text:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r, size=10.25)


def add_signature_block(doc, role):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(role)
    set_run_font(r, size=11, bold=True, color=NAVY)
    signature = add_label_line(doc, "Signature", "________________________________________________________")
    signature.paragraph_format.keep_with_next = True
    printed_name = add_label_line(doc, "Printed name", "___________________________________________________")
    printed_name.paragraph_format.keep_with_next = True
    if role == "Client Authorized Signer":
        business_name = add_label_line(doc, "Business legal name", "______________________________________________")
        business_name.paragraph_format.keep_with_next = True
    add_label_line(doc, "Date", "_____________________________________________________________")


def build_agreement(package):
    doc = Document()
    configure_document(doc, package["name"])

    title = doc.add_paragraph(style="Title")
    title.add_run(f"{package['name']} Monthly Services Agreement")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    r = subtitle.add_run(f"{package['scope']}  |  {package['price']} per month")
    set_run_font(r, size=12, bold=True, color=NAVY)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(12)
    r = intro.add_run(
        "This Agreement explains the monthly services Wolin will provide, the Client's responsibilities, "
        "the payment and cancellation terms, and the limits on results that any search or visibility service can promise. "
        "By signing, the Client confirms that the signer is authorized to enter this Agreement for the business named below."
    )
    set_run_font(r, size=10.75)

    add_heading(doc, "1 Parties and Effective Date")
    add_paragraph(doc, " Caleb Wolin operating under the name Wolin, referred to as Provider.", "Provider")
    add_paragraph(doc, " The business identified below, referred to as Client.", "Client")
    add_paragraph(doc, " This Agreement is effective on the date the Client signs it below.", "Effective date")

    add_heading(doc, "2 Package and Monthly Fee")
    add_paragraph(
        doc,
        f" Client selects the {package['name']} package at {package['price']} per monthly service period, plus any applicable tax. "
        "The package is month-to-month and does not require a fixed long-term commitment.",
        "Selected service"
    )
    add_paragraph(doc, f" {package['description']}", "Package purpose")
    add_paragraph(
        doc,
        " Each service period begins on the first payment date and renews monthly on the corresponding calendar date. "
        "If a month does not contain that date, renewal occurs on the month's last day.",
        "Service period"
    )

    add_heading(doc, "3 Services Included")
    add_paragraph(
        doc,
        " Provider will use reasonable professional judgment to prioritize and perform the following services during each monthly service period. "
        "The exact order and timing depend on account access, platform availability, Client approvals, current conditions, and the work with the greatest expected value.",
        "Scope"
    )
    add_bullets(doc, package["included"])

    add_heading(doc, "4 Service Schedule and Delivery")
    add_paragraph(
        doc,
        " Work begins after Provider receives the signed Agreement, the first monthly payment, and the access and materials reasonably needed for the work. "
        "Provider may communicate findings, completed changes, and next steps by email, shared document, report, or another format agreed with the Client.",
        "Start of work"
    )
    add_paragraph(
        doc,
        " The monthly fee reserves recurring professional time and covers the included services that Provider can reasonably complete during that service period. "
        "It does not purchase unlimited labor or require every possible recommendation to be completed in one month.",
        "Monthly capacity"
    )

    add_heading(doc, "5 Client Responsibilities")
    add_bullets(doc, [
        "Provide timely and lawful access to relevant profiles, website systems, analytics, hosting, domains, and third-party accounts",
        "Provide accurate business information, approved branding, images, service details, locations, hours, prices, policies, and claims",
        "Respond to approval requests and questions within a reasonable time and identify any required internal or legal review",
        "Maintain backups, licenses, subscriptions, hosting, domains, and security controls unless a written add-on states otherwise",
        "Tell Provider promptly about account notices, platform warnings, business changes, website changes, or work by other vendors that may affect the services",
        "Confirm that the Client has the right to use all materials and account access supplied to Provider",
    ])
    add_paragraph(
        doc,
        " Provider may pause affected work when access, information, approval, payment, or a safe backup is unavailable. "
        "A Client delay does not extend the paid service period unless Provider agrees in writing.",
        "Delays"
    )

    add_heading(doc, "6 Services Not Included")
    add_paragraph(
        doc,
        " The following are outside this package unless the parties approve a separate written add-on, statement of work, or price.",
        "Excluded work"
    )
    add_bullets(doc, package["excluded"])
    add_paragraph(
        doc,
        " Provider will request approval before purchasing a third-party product or service for the Client. Approved third-party costs are paid by the Client and are separate from the monthly fee.",
        "Outside costs"
    )

    add_heading(doc, "7 Results and Platform Limitations")
    add_paragraph(
        doc,
        " Provider promises to perform the included services with reasonable care and professional diligence. Provider does not promise a particular business result.",
        "Service standard"
    )
    add_bullets(doc, [
        "No specific search ranking, map position, indexing outcome, traffic level, impression count, click volume, review volume, lead volume, sale, revenue, or return on investment is guaranteed",
        "No mention, citation, recommendation, factual answer, placement, or frequency in ChatGPT, Claude, Google AI, or another AI system is guaranteed",
        "No platform verification, reinstatement, approval, publication, feature availability, or continued account access is guaranteed",
        "No particular Core Web Vitals score, loading time, accessibility score, conversion rate, or compatibility with every device, browser, theme, plugin, or third-party system is guaranteed",
        "Results may change because of competition, location, seasonality, demand, prior account history, Client decisions, algorithm changes, platform policies, outages, and third-party conduct",
    ])
    add_bullets(doc, package["specific_limits"])
    add_paragraph(
        doc,
        " Provider will not use deceptive practices, fake reviews, false business information, hidden text, link schemes, impersonation, or other tactics that Provider reasonably believes violate law or platform rules.",
        "Compliance"
    )

    add_heading(doc, "8 Fees and Payment")
    add_paragraph(
        doc,
        f" The Client will pay {package['price']} in advance for each monthly service period through Provider's payment link or another invoicing method approved by Provider. "
        "The first payment is due after signature and before work begins.",
        "Monthly fee"
    )
    add_paragraph(
        doc,
        " Provider may pause work while an amount is overdue. A pause for nonpayment does not waive the amount owed for work already performed or for the active service period. "
        "Provider will not increase the package fee without advance written notice. The Client may cancel before a new price takes effect.",
        "Payment problems and price changes"
    )

    add_heading(doc, "9 Cancellation and Ending Service")
    add_paragraph(
        doc,
        " The Client may cancel at any time by emailing caleb.wolin@gmail.com. Cancellation stops future monthly renewals and becomes effective at the end of the current paid service period. "
        "The Client should send the cancellation before the next renewal payment is due.",
        "Client cancellation"
    )
    add_paragraph(
        doc,
        " Monthly fees are not prorated or refunded for a service period that has already begun, except when required by law or when Provider agrees otherwise in writing. "
        "Provider will complete or reasonably wind down the work scheduled for the paid period, subject to Client access and cooperation.",
        "Current service period"
    )
    add_paragraph(
        doc,
        " Provider may end this Agreement on seven days' written notice. Provider may pause or end service immediately for nonpayment, unlawful instructions, abusive conduct, material breach, security risk, or loss of access needed to perform the work. "
        "If Provider ends service without cause before the paid period ends, Provider will refund the unused portion of that period on a reasonable prorated basis.",
        "Provider cancellation"
    )
    add_paragraph(
        doc,
        " After service ends, Provider will return or remove Client account access under Provider's control within a reasonable time. The Client remains responsible for changing shared passwords, removing user access, and maintaining its accounts and data.",
        "Account transition"
    )

    add_heading(doc, "10 Changes and Additional Work")
    add_paragraph(
        doc,
        " Work outside the included services requires the Client's written approval of the added scope, timing, and price. Email approval is sufficient. "
        "Provider is not required to begin added work before receiving any requested deposit or payment.",
        "Scope changes"
    )

    add_heading(doc, "11 Ownership and Use of Materials")
    add_paragraph(
        doc,
        " The Client keeps ownership of its pre-existing names, logos, photos, content, accounts, data, and other materials. After full payment, the Client owns final written content, configuration changes, and other deliverables created specifically for the Client under this Agreement, to the extent those items can be owned and transferred.",
        "Client materials and deliverables"
    )
    add_paragraph(
        doc,
        " Provider keeps ownership of pre-existing methods, templates, checklists, tools, code, processes, know-how, and general skills. To the extent a Provider-owned element is embedded in a paid deliverable, Provider grants the Client a perpetual, nonexclusive license to use that element as part of the deliverable for the Client's business.",
        "Provider materials"
    )
    add_paragraph(
        doc,
        " Third-party materials remain subject to their own licenses and terms. Provider does not transfer ownership of search platforms, AI systems, themes, plugins, stock media, fonts, or other third-party property.",
        "Third-party materials"
    )

    add_heading(doc, "12 Confidentiality and Account Access")
    add_paragraph(
        doc,
        " Each party will use reasonable care to protect nonpublic business information received from the other and will use it only to perform or receive the services. This obligation does not apply to information that is public through no breach, already lawfully known, independently developed, or lawfully received from another source.",
        "Confidential information"
    )
    add_paragraph(
        doc,
        " Provider may use employees, contractors, and service providers who reasonably need access to perform the work and who are expected to protect confidential information. The Client authorizes Provider to access and make changes in the accounts and systems supplied for the included services.",
        "Authorized access"
    )

    add_heading(doc, "13 Independent Contractor")
    add_paragraph(
        doc,
        " Provider is an independent contractor, not the Client's employee, partner, joint venturer, agent, or legal representative. Provider controls the manner and means of performing the services, subject to the agreed scope, applicable law, and platform rules.",
        "Relationship"
    )

    add_heading(doc, "14 Warranties and Disclaimers")
    add_paragraph(
        doc,
        " Each party represents that it has authority to enter this Agreement. The Client represents that information and instructions it supplies are accurate and lawful. Provider represents that Provider will perform the services in a professional manner consistent with ordinary industry practice.",
        "Mutual promises"
    )
    add_paragraph(
        doc,
        " Except for the express promises in this Agreement, the services and deliverables are provided as is and as available. To the extent permitted by law, Provider disclaims implied warranties, including merchantability, fitness for a particular purpose, and noninfringement. Nothing in this Agreement excludes a warranty or right that cannot lawfully be excluded.",
        "Other warranties"
    )

    add_heading(doc, "15 Limitation of Liability")
    add_paragraph(
        doc,
        " To the extent permitted by law, neither party is liable to the other for indirect, incidental, special, exemplary, punitive, or consequential damages, or for lost profits, lost revenue, lost opportunities, loss of goodwill, or loss of data arising from this Agreement, even if advised that such loss is possible.",
        "Excluded losses"
    )
    add_paragraph(
        doc,
        " To the extent permitted by law, Provider's total liability arising from this Agreement will not exceed the fees the Client paid Provider under this Agreement during the three months immediately before the event giving rise to the claim. This limit does not apply where a limit is prohibited by law or to Provider's fraud, willful misconduct, or gross negligence.",
        "Liability cap"
    )

    add_heading(doc, "16 General Terms")
    add_paragraph(
        doc,
        " The parties will first try in good faith to resolve a dispute through direct discussion. Either party may seek available legal remedies if the dispute is not resolved. The laws of the state where Provider principally resides on the Effective Date govern this Agreement, without regard to conflict-of-law rules, unless the parties write a different governing state here.",
        "Disputes and governing law"
    )
    add_label_line(doc, "Different governing state if agreed", "________________________________________")
    add_paragraph(
        doc,
        " This Agreement and any written add-on are the entire agreement about these services and replace prior discussions or proposals on the same subject. A change must be in writing and accepted by both parties. If one provision is unenforceable, the rest remains effective. A failure to enforce a provision is not a waiver. Neither party may assign this Agreement without the other's written consent, except as part of a sale or transfer of substantially all of that party's relevant business or assets.",
        "Entire agreement"
    )
    add_paragraph(
        doc,
        " Notices may be sent by email to the addresses the parties use for service communications. The Client must send cancellation notices to caleb.wolin@gmail.com. Signatures in counterparts and electronic signatures are effective, and a scanned or electronic copy may be treated as an original.",
        "Notices and signatures"
    )

    add_heading(doc, "17 Client Signature")
    signature_agreement = add_paragraph(
        doc,
        f" By signing below, the Client acknowledges that it has read and agrees to the {package['name']} package at {package['price']} per month and to all terms in this Agreement.",
        "Agreement"
    )
    signature_agreement.paragraph_format.keep_with_next = True
    add_signature_block(doc, "Client Authorized Signer")

    output = OUTPUT_DIR / f"Wolin_{package['name']}_Monthly_Services_Agreement.docx"
    doc.save(output)
    return output


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    outputs = [build_agreement(package) for package in PACKAGES]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
